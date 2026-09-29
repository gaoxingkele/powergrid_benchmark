"""Author-review Word copy from paper_information.tex (MDPI Information).

Embeds figures at 300 dpi and numbered display equations as rendered images.
SuSy should still take the LaTeX ZIP + PDF; this file is for author review.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
TEX = HERE / "paper_information.tex"
OUT = HERE / "paper_information.docx"
FIG_DIR = HERE / "figures"
CACHE = HERE / "_docx_assets"


def brace(src: str, start: int) -> tuple[str, int]:
    assert src[start] == "{"
    depth = 0
    i = start
    while i < len(src):
        if src[i] == "{":
            depth += 1
        elif src[i] == "}":
            depth -= 1
            if depth == 0:
                return src[start + 1 : i], i + 1
        i += 1
    raise ValueError("unbalanced")


def cmd(src: str, name: str) -> str:
    token = "\\" + name + "{"
    i = src.find(token)
    if i < 0:
        return ""
    text, _ = brace(src, i + len("\\" + name))
    return text


def cmd2(src: str, name: str) -> tuple[str, str]:
    token = "\\" + name + "{"
    i = src.find(token)
    if i < 0:
        return "", ""
    first, j = brace(src, i + len("\\" + name))
    while j < len(src) and src[j].isspace():
        j += 1
    if j < len(src) and src[j] == "{":
        second, _ = brace(src, j)
        return first, second
    return first, ""


def strip_tex(text: str) -> str:
    text = text.replace("\\cges{}", "C²GES")
    text = text.replace("\\LaTeX{}", "LaTeX")
    text = re.sub(
        r"\\texorpdfstring\{([^{}]*)\}\{([^{}]*)\}",
        r"\1",
        text,
    )
    text = text.replace("\\textsuperscript{2}", "²")
    text = text.replace("\\&", "&")
    text = text.replace("\\%", "%")
    text = text.replace("\\_", "_")
    text = text.replace("\\#", "#")
    text = text.replace("~", " ")
    text = text.replace("\\,", " ")
    text = text.replace("\\;", " ")
    text = text.replace("\\nolinkurl", "")
    text = text.replace("\\noindent", "")
    text = text.replace("\\centering", "")
    text = text.replace("\\scriptsize", "")
    text = text.replace("\\small", "")
    text = text.replace("\\textit", "")
    text = re.sub(r"\\(emph|textbf|texttt|textrm)\{([^{}]*)\}", r"\2", text)
    text = re.sub(r"\\href\{[^{}]*\}\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\cite\{[^}]*\}", "", text)
    text = re.sub(r"\\(?:C)?ref\{[^}]*\}", "", text)
    text = re.sub(r"\\eqref\{[^}]*\}", "", text)
    text = re.sub(r"\\label\{[^}]*\}", "", text)
    text = re.sub(r"\$([^$]+)\$", r"\1", text)
    text = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    text = text.replace("  ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def set_run_font(run, name="Times New Roman", size=11, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_p(doc, text, *, size=11, bold=False, italic=False, center=False, space_after=8, first_line=None):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if first_line is not None:
        p.paragraph_format.first_line_indent = Inches(first_line)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def pdf_to_png(pdf_path: Path, png_path: Path, dpi: int = 300) -> Path:
    png_path.parent.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(pdf_path)
    page = doc[0]
    pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
    pix.save(png_path)
    doc.close()
    return png_path


def render_equation(latex: str, index: int) -> Path:
    CACHE.mkdir(parents=True, exist_ok=True)
    png = CACHE / f"eq_{index:02d}.png"
    work = Path(tempfile.mkdtemp(prefix="c2ges_eq_"))
    try:
        body = latex.strip().rstrip(".")
        src = r"""
\documentclass{article}
\usepackage{amsmath,amssymb,amsfonts}
\usepackage[active,tightpage]{preview}
\pagestyle{empty}
\begin{document}
\begin{preview}
\begin{equation*}
%s
\end{equation*}
\end{preview}
\end{document}
""" % body
        tex_path = work / "eq.tex"
        tex_path.write_text(src, encoding="utf-8")
        subprocess.run(
            ["pdflatex", "-interaction=nonstopmode", "-halt-on-error", "eq.tex"],
            cwd=work,
            check=True,
            capture_output=True,
        )
        pdf_to_png(work / "eq.pdf", png, dpi=300)
    finally:
        shutil.rmtree(work, ignore_errors=True)
    return png


def add_picture(doc, image: Path, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(image), width=Inches(width))
    return p


def add_table(doc, caption: str, rows: list[list[str]]):
    if caption:
        add_p(doc, caption, size=10, italic=True, center=True, space_after=4)
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = row[j] if j < len(row) else ""
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_font(run, size=9, bold=(i == 0))
    doc.add_paragraph()


def parse_tabular(block: str) -> list[list[str]]:
    block = re.sub(r"\\(toprule|midrule|bottomrule|hline)", r"\\\\", block)
    rows = []
    for raw in re.split(r"\\\\", block):
        raw = raw.strip()
        if not raw or raw.startswith("\\") and "&" not in raw:
            continue
        cells = [strip_tex(c) for c in raw.split("&")]
        if any(cells):
            rows.append(cells)
    return rows


def extract_env(src: str, start: int, env: str) -> tuple[str, int]:
    open_tok = "\\begin{" + env + "}"
    close_tok = "\\end{" + env + "}"
    assert src.startswith(open_tok, start)
    i = start + len(open_tok)
    depth = 1
    while i < len(src):
        nxt_open = src.find(open_tok, i)
        nxt_close = src.find(close_tok, i)
        if nxt_close < 0:
            raise ValueError("unclosed " + env)
        if nxt_open >= 0 and nxt_open < nxt_close:
            depth += 1
            i = nxt_open + len(open_tok)
            continue
        depth -= 1
        if depth == 0:
            return src[start + len(open_tok) : nxt_close], nxt_close + len(close_tok)
        i = nxt_close + len(close_tok)
    raise ValueError("unclosed " + env)


def skip_optional(src: str, i: int) -> int:
    if i < len(src) and src[i] == "[":
        j = src.find("]", i)
        return j + 1 if j >= 0 else i
    return i


FIGURE_WIDTH = {
    "fig01_algorithm_dual_panel.pdf": 6.3,
    "fig02_dataset_flow.pdf": 6.3,
    "fig03_aggregate_rougel.pdf": 6.3,
    "fig04_paired_differences.pdf": 6.3,
    "fig05_output_length.pdf": 5.8,
    "fig06_component_diagnostic.pdf": 6.3,
}


def emit_body(doc: Document, body: str) -> None:
    i = 0
    n = len(body)
    eq_index = 0
    fig_index = 0
    tab_index = 0
    buf: list[str] = []

    def flush():
        text = strip_tex("".join(buf))
        buf.clear()
        if not text:
            return
        for para in re.split(r"\n\s*\n", text):
            para = re.sub(r"\s+", " ", para).strip()
            if para:
                add_p(doc, para, size=11, space_after=8, first_line=0.3)

    while i < n:
        if body.startswith("\\section{", i):
            flush()
            heading, i = brace(body, i + len("\\section"))
            add_p(doc, strip_tex(heading), size=13, bold=True, space_after=8)
            continue
        if body.startswith("\\subsection{", i):
            flush()
            heading, i = brace(body, i + len("\\subsection"))
            add_p(doc, strip_tex(heading), size=12, bold=True, space_after=6)
            continue
        if body.startswith("\\subsubsection{", i):
            flush()
            heading, i = brace(body, i + len("\\subsubsection"))
            add_p(doc, strip_tex(heading), size=11, bold=True, italic=True, space_after=6)
            continue
        if body.startswith("\\begin{figure}", i):
            flush()
            fig_index += 1
            block, i = extract_env(body, i, "figure")
            inc = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}", block)
            cap = re.search(r"\\caption\{", block)
            caption = ""
            if cap:
                caption, _ = brace(block, cap.start() + len("\\caption"))
            if inc:
                pdf = HERE / inc.group(1)
                png = CACHE / (pdf.stem + ".png")
                pdf_to_png(pdf, png, dpi=300)
                add_picture(doc, png, width=FIGURE_WIDTH.get(pdf.name, 6.3))
            add_p(
                doc,
                f"Figure {fig_index}. " + strip_tex(caption),
                size=10,
                italic=True,
                center=True,
                space_after=10,
            )
            continue
        if body.startswith("\\begin{table}", i):
            flush()
            tab_index += 1
            block, i = extract_env(body, i, "table")
            cap = re.search(r"\\caption\{", block)
            caption = ""
            if cap:
                caption, _ = brace(block, cap.start() + len("\\caption"))
            tab = re.search(r"\\begin\{(tabularx|tabular)\}", block)
            rows: list[list[str]] = []
            if tab:
                env = tab.group(1)
                inner, _ = extract_env(block, tab.start(), env)
                # drop column spec
                inner = re.sub(r"^\{[^{}]*\}", "", inner.strip(), count=1)
                inner = re.sub(r"^\{[^{}]*\}", "", inner.strip(), count=1)
                rows = parse_tabular(inner)
            add_table(doc, f"Table {tab_index}. " + strip_tex(caption), rows)
            continue
        if body.startswith("\\begin{equation}", i):
            flush()
            eq_index += 1
            block, i = extract_env(body, i, "equation")
            math = re.sub(r"\\label\{[^}]*\}", "", block).strip()
            try:
                png = render_equation(math, eq_index)
                add_picture(doc, png, width=5.6)
            except Exception as exc:
                add_p(doc, strip_tex(math), size=11, italic=True, center=True, space_after=2)
                add_p(doc, f"[equation render failed: {exc}]", size=8, center=True, space_after=2)
            add_p(doc, f"Equation ({eq_index})", size=10, italic=True, center=True, space_after=10)
            continue
        if body.startswith("\\begin{itemize}", i) or body.startswith("\\begin{enumerate}", i):
            env = "itemize" if body.startswith("\\begin{itemize}", i) else "enumerate"
            flush()
            block, i = extract_env(body, i, env)
            for item in re.split(r"\\item\b", block)[1:]:
                add_p(doc, "• " + strip_tex(item), size=11, space_after=4)
            continue
        buf.append(body[i])
        i += 1
    flush()


def main() -> None:
    src = TEX.read_text(encoding="utf-8")
    title = strip_tex(cmd(src, "Title")).replace("C2GES", "C²GES")
    if title.startswith("C"):
        title = title.replace("CGES", "C²GES")
    title = re.sub(r"^C\s*2\s*GES|^C²GES|^CGES", "C²GES", title)
    # Title command uses \texorpdfstring
    raw_title = cmd(src, "Title")
    title = strip_tex(raw_title.replace(r"\texorpdfstring{\textsuperscript{2}}{2}", "²"))
    authors = strip_tex(cmd(src, "AuthorNames"))
    abstract = strip_tex(cmd(src, "abstract"))
    keywords = strip_tex(cmd(src, "keyword"))
    corres = strip_tex(cmd(src, "corres"))
    address = (
        "1 NARI Group Corporation (State Grid Electric Power Research Institute), "
        "Nanjing 211106, China; 2 Beijing Kedong Electric Power Control System Co., Ltd., "
        "Beijing 100080, China"
    )
    body = src.split("\\begin{document}", 1)[1]
    body = body.split("\\supplementary", 1)[0]
    declarations = {
        "Supplementary Materials": strip_tex(cmd(src, "supplementary")),
        "Author Contributions": strip_tex(cmd(src, "authorcontributions")),
        "Funding": strip_tex(cmd(src, "funding")),
        "Institutional Review Board Statement": strip_tex(cmd(src, "institutionalreview")),
        "Informed Consent Statement": strip_tex(cmd(src, "informedconsent")),
        "Data Availability Statement": strip_tex(cmd(src, "dataavailability")),
        "Acknowledgments": strip_tex(cmd(src, "acknowledgments")),
        "Conflicts of Interest": strip_tex(cmd(src, "conflictsofinterest")),
        "Abbreviations": (
            "The following abbreviations are used in this manuscript. "
            "C²GES: role-conditioned extractive selector described in this paper; "
            "ENTSO-E: European Network of Transmission System Operators for Electricity; "
            "LOSO: leave-one-series-out; "
            "MMR: maximal marginal relevance; "
            "NERC: North American Electric Reliability Corporation; "
            "NESO: National Energy System Operator; "
            "ROUGE: Recall-Oriented Understudy for Gisting Evaluation; "
            "ROUGE-L: ROUGE longest-common-subsequence F1."
        ),
    }

    doc = Document()
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    add_p(doc, "Article submitted to Information (MDPI)", size=11, italic=True, center=True, space_after=4)
    add_p(doc, title, size=16, bold=True, center=True, space_after=8)
    add_p(doc, authors, size=12, bold=True, center=True, space_after=4)
    add_p(doc, address, size=10, center=True, space_after=2)
    add_p(doc, corres, size=10, center=True, space_after=12)
    add_p(doc, "Abstract", size=12, bold=True, space_after=4)
    add_p(doc, abstract, size=11, italic=True, space_after=8)
    add_p(doc, "Keywords: " + keywords.replace(";", "; "), size=11, space_after=12)

    emit_body(doc, body)

    add_p(doc, "Declarations", size=13, bold=True, space_after=6)
    for name, text in declarations.items():
        if text:
            add_p(doc, name, size=12, bold=True, space_after=2)
            add_p(doc, text, size=11, space_after=8)

    note = (
        "This Word file is an author-review conversion from the MDPI LaTeX source "
        "(class v6.5a, 11 September 2026). Display equations are rendered from the "
        "same LaTeX as the PDF; figures are embedded at 300 dpi. SuSy submission "
        "should use the LaTeX ZIP and the compiled PDF. This is a diagnostic, "
        "non-confirmatory study. It does not claim system superiority, "
        "structural-proxy validity, or operational benefit."
    )
    p = add_p(doc, note, size=9, space_after=0)
    p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.save(OUT)
    print(f"wrote {OUT} bytes={OUT.stat().st_size}")


if __name__ == "__main__":
    main()
