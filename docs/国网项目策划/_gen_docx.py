# -*- coding: utf-8 -*-
"""把三家单位项目策划书_草稿.md 转成规范 Word 文档"""
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC = r"D:\BaiduSyncdisk\2026-08\国网项目策划\三家单位项目策划书_草稿.md"
DST = r"D:\BaiduSyncdisk\2026-08\国网项目策划\三家单位电力AI项目策划书.docx"

def set_font(run, name_cn="仿宋_GB2312", name_en="Times New Roman", size=12, bold=False, color=None):
    run.font.name = name_en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name_cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)

doc = Document()
sec = doc.sections[0]
sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.6)
sec.top_margin = Cm(3.7); sec.bottom_margin = Cm(3.5)

# 封面标题
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("三家单位电力人工智能项目策划书"), "方正小标宋简体", size=22, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("(湖北电科院 · 河南周口供电公司 · 冀北智能配网中心)"), "楷体_GB2312", size=15)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("2026年8月"), "仿宋_GB2312", size=14)
doc.add_paragraph()

with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

LABELS = ("**对应指南条款**", "**项目目标**", "**基本建设内容**", "**技术路线**", "**关键技术问题**")

for line in lines:
    s = line.strip()
    if not s or s.startswith(">"):
        continue
    if s.startswith("# ") :
        t = s[2:].strip()
        if t.startswith("三家单位"):  # 跳过 md 总标题(封面已有)
            continue
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(10)
        set_font(p.add_run(t), "黑体", size=16, bold=True)
    elif s.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
        set_font(p.add_run(s[3:].strip()), "黑体", size=14, bold=True)
    elif s.startswith(tuple(LABELS)):
        m = re.match(r"\*\*(.+?)\*\*[::]\s*(.*)", s)
        label, body = m.group(1), m.group(2)
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.85)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = Pt(24)
        set_font(p.add_run(f"【{label}】"), "黑体", size=12, bold=True)
        set_font(p.add_run(body), "仿宋_GB2312", size=12)
    else:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.85)
        p.paragraph_format.line_spacing = Pt(24)
        set_font(p.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", s)), "仿宋_GB2312", size=12)

doc.save(DST)
print("saved:", DST)
