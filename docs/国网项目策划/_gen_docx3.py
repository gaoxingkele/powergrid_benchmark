# -*- coding: utf-8 -*-
"""按单位拆分为三个 Word 文档:
- 【对应指南条款】置于每个专题尾部,斜体不加粗
- 建设内容/技术路线/关键技术问题中"一是二是三是"式内容转编号列表
"""
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC = r"D:\BaiduSyncdisk\2026-08\国网项目策划\三家单位项目策划书_草稿.md"

UNITS = [
    ("第一部分", "湖北电科院", r"D:\BaiduSyncdisk\2026-08\国网项目策划\湖北电科院电力AI项目策划书.docx"),
    ("第二部分", "河南周口供电公司", r"D:\BaiduSyncdisk\2026-08\国网项目策划\周口供电公司电力AI项目策划书.docx"),
    ("第三部分", "冀北智能配网中心", r"D:\BaiduSyncdisk\2026-08\国网项目策划\冀北智能配网中心电力AI项目策划书.docx"),
]

def set_font(run, cn="仿宋_GB2312", en="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic

CN_NUM = "一二三四五六七八九"

def split_enum(text):
    """把 '一是…;二是…' 拆为 (preamble, [items]);不含枚举则返回 (text, [])"""
    if "一是" not in text or "二是" not in text:
        return text, []
    parts = re.split(r"[;;]\s*(?=[%s]是)" % CN_NUM, text)
    preamble = ""
    if not re.match(r"^[%s]是" % CN_NUM, parts[0]):
        preamble = parts[0].rstrip(";;")
        parts = parts[1:]
    items = [re.sub(r"^[%s]是\s*[,,]?\s*" % CN_NUM, "", p).strip() for p in parts if p.strip()]
    return preamble, items

# ---------- 解析 markdown ----------
with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

parts = {}          # part_key -> [topic, ...]
cur_part = None
cur_topic = None
part_note = {}      # part 标题行附注(如"均联合省电科院申报")

for line in lines:
    s = line.strip()
    if s.startswith("# 第"):
        for key, _, _ in UNITS:
            if key in s:
                cur_part = key
                parts[cur_part] = []
                m = re.search(r"[((](.+?)[))]\s*$", s)
                part_note[cur_part] = m.group(1) if m else ""
        continue
    if cur_part is None:
        continue
    if s.startswith("## "):
        cur_topic = {"title": s[3:].strip(), "对应指南条款": "", "项目目标": "",
                     "基本建设内容": "", "技术路线": "", "关键技术问题": ""}
        parts[cur_part].append(cur_topic)
        continue
    if cur_topic is None or not s:
        continue
    m = re.match(r"\*\*(对应指南条款|项目目标|基本建设内容|技术路线|关键技术问题)\*\*[::]\s*(.*)", s)
    if m:
        body = re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2))
        cur_topic[m.group(1)] = body

# ---------- 生成三个文档 ----------
for key, unit, dst in parts.keys() and [(k, u, d) for k, u, d in UNITS]:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.6)
    sec.top_margin = Cm(3.7); sec.bottom_margin = Cm(3.5)

    # 封面
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"{unit}电力人工智能项目策划书"), "方正小标宋简体", size=22, bold=True)
    note = part_note.get(key, "")
    if note:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(f"({note})"), "楷体_GB2312", size=14)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("2026年8月"), "仿宋_GB2312", size=14)
    doc.add_paragraph()

    for topic in parts[key]:
        # 专题标题
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
        set_font(p.add_run(topic["title"]), "黑体", size=14, bold=True)

        # 项目目标:标签+正文同段
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.85)
        p.paragraph_format.line_spacing = Pt(24); p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run("【项目目标】"), "黑体", size=12, bold=True)
        set_font(p.add_run(topic["项目目标"]), "仿宋_GB2312", size=12)

        # 建设内容/技术路线/关键技术问题:枚举转列表
        for label in ("基本建设内容", "技术路线", "关键技术问题"):
            body = topic[label]
            preamble, items = split_enum(body)
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.85)
            p.paragraph_format.line_spacing = Pt(24); p.paragraph_format.space_after = Pt(4)
            set_font(p.add_run(f"【{label}】"), "黑体", size=12, bold=True)
            if items:
                if preamble:
                    set_font(p.add_run(preamble + "。"), "仿宋_GB2312", size=12)
                for i, it in enumerate(items, 1):
                    lp = doc.add_paragraph()
                    lp.paragraph_format.left_indent = Cm(1.2)
                    lp.paragraph_format.first_line_indent = Cm(0)
                    lp.paragraph_format.line_spacing = Pt(24); lp.paragraph_format.space_after = Pt(2)
                    if not it.endswith(("。", ";", ";")):
                        it += "。"
                    it = it.rstrip(";;") + ("。" if not it.endswith("。") else "")
                    set_font(lp.add_run(f"({CN_NUM[i-1]}) {it}"), "仿宋_GB2312", size=12)
            else:
                set_font(p.add_run(body), "仿宋_GB2312", size=12)

        # 对应指南条款:尾部,斜体不加粗
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.85)
        p.paragraph_format.line_spacing = Pt(24); p.paragraph_format.space_after = Pt(10)
        set_font(p.add_run("对应指南条款:" + topic["对应指南条款"]), "楷体_GB2312", size=12, italic=True)

    doc.save(dst)
    print("saved:", dst, f"({len(parts[key])}题)")
